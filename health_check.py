#!/usr/bin/env python3
from __future__ import annotations

import argparse
import re
from collections import Counter
from pathlib import Path
from typing import Dict, List, Optional, Set, Tuple

from docx import Document

from config import canonical_okved, load_table_source_map
from mo import canonical_mo
from table_manager import TableManager

TAG_REGEX = re.compile(r"\{\{([^}]+)\}\}")
YEAR_RE = re.compile(r"^\d{2}$|^20\d{2}$")


def _map_year_to_suffix(year: str) -> Optional[str]:
    if not year or not year.isdigit():
        return None
    if len(year) == 4 and year.startswith("20"):
        return year[-2:]
    if len(year) == 2:
        return year
    return None


def extract_tags_from_docx(docx_path: Path) -> List[str]:
    doc = Document(docx_path)
    tags: List[str] = []

    for p in doc.paragraphs:
        tags.extend(TAG_REGEX.findall(p.text or ""))

    tm = TableManager(doc)
    for table in tm.iter_tables():
        for row in table.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    tags.extend(TAG_REGEX.findall(p.text or ""))

    return tags


def parse_tag_to_keys(full_tag: str, known_entities: Set[str]) -> Tuple[List[str], Optional[str], Optional[str]]:
    """
    Returns:
      - candidate keys in master_data style: ["ENTITY|indicator_yy", ...]
      - indicator base name
      - source prefix (OKVED/MO/None)
    """
    raw = full_tag
    source_prefix = None
    if raw.startswith("OKVED_"):
        source_prefix = "OKVED"
        raw = raw[len("OKVED_"):]
    elif raw.startswith("MO_"):
        source_prefix = "MO"
        raw = raw[len("MO_"):]

    parts = raw.split("_")
    if len(parts) < 2:
        return [], None, source_prefix

    # robust split by known entity keys first (supports compound entity code)
    entity_code = None
    indicator_parts: List[str] = []

    for split_idx in range(len(parts) - 1, 0, -1):
        entity_candidate_raw = "_".join(parts[:split_idx])
        if source_prefix == "MO":
            # Префикс "MO_" в ключе master_data — см. data_filler_v2.py,
            # _infer_entity_key: без него код МО может случайно совпасть с
            # чужим ключом в общем плоском словаре.
            entity_candidate = f"MO_{canonical_mo(entity_candidate_raw)}"
        else:
            if "." in entity_candidate_raw or any(c.isalpha() for c in entity_candidate_raw):
                entity_candidate = canonical_okved(entity_candidate_raw.replace("_", "."))
            else:
                entity_candidate = canonical_okved(entity_candidate_raw)

        if entity_candidate in known_entities:
            entity_code = entity_candidate
            indicator_parts = parts[split_idx:]
            break

    # fallback naive split
    if entity_code is None:
        entity_raw = "_".join(parts[:-1])
        if source_prefix == "MO":
            entity_code = f"MO_{canonical_mo(entity_raw)}"
        else:
            if "." in entity_raw or any(c.isalpha() for c in entity_raw):
                entity_code = canonical_okved(entity_raw.replace("_", "."))
            else:
                entity_code = canonical_okved(entity_raw)
        indicator_parts = [parts[-1]]

    if not indicator_parts:
        return [], None, source_prefix

    year_suffix = None
    if YEAR_RE.fullmatch(indicator_parts[-1]):
        year_suffix = _map_year_to_suffix(indicator_parts[-1])
        if year_suffix:
            indicator_parts = indicator_parts[:-1]

    if not indicator_parts:
        return [], None, source_prefix

    indicator = "_".join(indicator_parts)

    keys: List[str] = []
    if year_suffix:
        keys.append(f"{entity_code}|{indicator}_{year_suffix}")
        # same fallback order as fill logic
        if year_suffix != "23":
            keys.append(f"{entity_code}|{indicator}_23")
        if year_suffix != "22":
            keys.append(f"{entity_code}|{indicator}_22")
    else:
        keys.append(f"{entity_code}|{indicator}")
        keys.append(f"{entity_code}|{indicator}_22")
        keys.append(f"{entity_code}|{indicator}_23")

    return keys, indicator, source_prefix


def build_master_key_set(master_data: Dict) -> Set[str]:
    keys = set()
    for entity, indicators in master_data.items():
        for ind_key in indicators.keys():
            keys.add(f"{entity}|{ind_key}")
    return keys


def infer_source_file_for_tag(tag: str, table_map: Dict[str, str]) -> str:
    # heuristic only: decide MO vs OKVED scope
    if tag.startswith("MO_"):
        mo_files = sorted({v for v in table_map.values() if "mo" in v.lower()})
        return mo_files[0] if mo_files else "<unknown_mo_source>"
    okved_files = sorted({v for v in table_map.values() if "mo" not in v.lower()})
    return okved_files[0] if okved_files else "<unknown_okved_source>"


def main() -> int:
    parser = argparse.ArgumentParser(description="Health check for DOCX tag -> master_data coverage")
    parser.add_argument("--template", type=Path, default=Path("output/Бюллетень_ШАБЛОН_С_ТЕГАМИ_v2.docx"))
    parser.add_argument("--top-missing", type=int, default=50)
    parser.add_argument("--top-sources", type=int, default=20)
    args = parser.parse_args()

    # Import late to avoid heavy deps unless script runs
    from data_filler_v2 import pre_load_all_excel_data_v2
    from config import load_okved_map
    from config_v2 import load_column_mapping_v2, validate_table_source_mapping

    base_dir = Path(__file__).parent.resolve()
    input_dir = base_dir / "input"
    mappings_dir = input_dir / "mappings"
    excel_dir = input_dir / "excel"

    okved_file = mappings_dir / "okved_mapping.csv"
    mo_file = mappings_dir / "mo.csv"
    table_mapping_file = mappings_dir / "table_source_data_mapping.csv"
    column_mapping_file = mappings_dir / "column_mapping_v2.csv"

    okved_to_name, name_to_okved_cleaned = load_okved_map(okved_file)
    table_map = load_table_source_map(table_mapping_file)
    validate_table_source_mapping(table_map, excel_dir)
    load_column_mapping_v2(column_mapping_file)

    master_data, _ = pre_load_all_excel_data_v2(
        excel_dir=excel_dir,
        table_source_mapping=table_map,
        okved_codes_set=set(okved_to_name.keys()),
        okved_name_to_code=name_to_okved_cleaned,
        column_mapping_path=column_mapping_file,
        mo_map_path=mo_file,
        use_fuzzy_match=True,
        fuzzy_threshold=0.80,
    )

    raw_tags = extract_tags_from_docx(args.template)
    unique_tags = sorted(set(raw_tags))

    master_keys = build_master_key_set(master_data)
    known_entities = set(master_data.keys())

    missing_tags: List[str] = []
    invalid_tags: List[str] = []
    matched = 0
    missing_indicators = Counter()
    source_issues = Counter()

    for t in unique_tags:
        candidate_keys, indicator, _ = parse_tag_to_keys(t, known_entities)
        if not candidate_keys:
            invalid_tags.append(t)
            source_issues[infer_source_file_for_tag(t, table_map)] += 1
            continue

        if any(k in master_keys for k in candidate_keys):
            matched += 1
        else:
            missing_tags.append(t)
            if indicator:
                missing_indicators[indicator] += 1
            source_issues[infer_source_file_for_tag(t, table_map)] += 1

    total = len(unique_tags)
    missing = len(missing_tags)
    invalid = len(invalid_tags)
    coverage = 0.0 if total == 0 else ((total - missing) / total) * 100.0

    print("\n===== HEALTH CHECK REPORT =====")
    print(f"Total unique tags: {total}")
    print(f"Matched tags: {matched}")
    print(f"Missing tags: {missing}")
    print(f"Invalid tags: {invalid}")
    print(f"Coverage: {coverage:.2f}%")

    print(f"\n--- TOP-{args.top_missing} MISSING INDICATORS ---")
    for ind, cnt in missing_indicators.most_common(args.top_missing):
        print(f"{ind}: {cnt}")

    print(f"\n--- TOP-{args.top_sources} PROBLEMATIC SOURCE FILES (heuristic) ---")
    for src, cnt in source_issues.most_common(args.top_sources):
        print(f"{src}: {cnt}")

    verdict = "FAIL" if missing > 0 else "PASS"
    print(f"\nVERDICT: {verdict}")

    return 1 if verdict == "FAIL" else 0


if __name__ == "__main__":
    raise SystemExit(main())