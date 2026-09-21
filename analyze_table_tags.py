#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
analyze_table_tags.py

Анализирует все ли таблицы в документе DOCX заполняются тегами.

Проверяет:
1. Количество таблиц в документе
2. Наличие тегов в каждой таблице
3. Какие таблицы НЕ содержат тегов (неполное покрытие)
4. Структуру каждой таблицы (количество строк, столбцов)
5. Соответствие между структурой и наличием тегов
"""

import re
from pathlib import Path
from typing import Dict, List, Tuple
from docx import Document


def analyze_document_structure(doc_path: str) -> Dict:
    """
    Анализирует структуру документа и наличие тегов в таблицах.
    
    Возвращает:
        {
            'total_tables': int,
            'tables': [
                {
                    'index': int,
                    'rows': int,
                    'cols': int,
                    'has_tags': bool,
                    'tag_count': int,
                    'tags_sample': list,
                    'cells_with_tags': int,
                    'empty_cells': int,
                    'cells_with_text': int,
                    'coverage_percent': float,
                    'issues': list
                },
                ...
            ],
            'summary': {
                'tables_with_tags': int,
                'tables_without_tags': int,
                'total_tags': int,
                'coverage_percent': float
            }
        }
    """
    
    doc = Document(doc_path)
    tag_regex = re.compile(r"\{\{([^}]+)\}\}")
    
    analysis = {
        'total_tables': len(doc.tables),
        'tables': [],
        'summary': {
            'tables_with_tags': 0,
            'tables_without_tags': 0,
            'total_tags': 0,
            'tables_fully_tagged': 0,
            'tables_partially_tagged': 0
        }
    }
    
    total_potential_cells = 0
    total_tagged_cells = 0
    
    for table_idx, table in enumerate(doc.tables):
        rows_count = len(table.rows)
        cols_count = len(table.columns) if table.columns else 0
        potential_cells = rows_count * cols_count if cols_count > 0 else 0
        
        table_analysis = {
            'index': table_idx + 1,
            'rows': rows_count,
            'cols': cols_count,
            'potential_cells': potential_cells,
            'has_tags': False,
            'tag_count': 0,
            'tags_sample': [],
            'cells_with_tags': 0,
            'cells_with_text': 0,
            'cells_with_dash': 0,
            'empty_cells': 0,
            'coverage_percent': 0.0,
            'issues': []
        }
        
        all_tags = []
        cells_with_tags = set()
        cells_with_content = {}
        
        # Анализируем все ячейки таблицы
        for row_idx, row in enumerate(table.rows):
            for col_idx, cell in enumerate(row.cells):
                cell_key = f"{row_idx},{col_idx}"
                cell_text = cell.text.strip()
                
                if not cell_text:
                    table_analysis['empty_cells'] += 1
                    cells_with_content[cell_key] = 'empty'
                elif cell_text == '-':
                    table_analysis['cells_with_dash'] += 1
                    cells_with_content[cell_key] = 'dash'
                elif cell_text == ' ':
                    cells_with_content[cell_key] = 'space'
                else:
                    table_analysis['cells_with_text'] += 1
                    cells_with_content[cell_key] = 'text'
                
                # Ищем теги
                tags = tag_regex.findall(cell_text)
                if tags:
                    all_tags.extend(tags)
                    cells_with_tags.add(cell_key)
                    table_analysis['cells_with_tags'] += 1
        
        # Заполняем результаты
        table_analysis['tag_count'] = len(all_tags)
        table_analysis['tags_sample'] = list(set(all_tags[:10]))  # Уникальные примеры
        table_analysis['has_tags'] = len(all_tags) > 0
        
        if potential_cells > 0:
            table_analysis['coverage_percent'] = round(
                (table_analysis['cells_with_tags'] / potential_cells) * 100, 2
            )
        
        # Определяем степень покрытия
        if table_analysis['cells_with_tags'] == 0:
            table_analysis['status'] = 'NOT_TAGGED'
            analysis['summary']['tables_without_tags'] += 1
        elif table_analysis['coverage_percent'] == 100:
            table_analysis['status'] = 'FULLY_TAGGED'
            analysis['summary']['tables_fully_tagged'] += 1
        elif table_analysis['coverage_percent'] > 0:
            table_analysis['status'] = 'PARTIALLY_TAGGED'
            analysis['summary']['tables_partially_tagged'] += 1
        
        if table_analysis['has_tags']:
            analysis['summary']['tables_with_tags'] += 1
        
        # Ищем проблемы
        if table_analysis['empty_cells'] == potential_cells:
            table_analysis['issues'].append('ПУСТА - все ячейки пусты')
        
        if table_analysis['rows'] < 2:
            table_analysis['issues'].append('ОЧЕНЬ_МАЛА - менее 2 строк')
        
        if table_analysis['cells_with_tags'] > 0 and table_analysis['coverage_percent'] < 50:
            table_analysis['issues'].append(f'СЛАБОЕ_ПОКРЫТИЕ - только {table_analysis["coverage_percent"]}%')
        
        if table_analysis['cells_with_dash'] > table_analysis['cells_with_tags']:
            table_analysis['issues'].append('БОЛЬШЕ_ПРОЧЕРКОВ_ЧЕМ_ТЕГОВ')
        
        analysis['tables'].append(table_analysis)
        
        # Суммируем
        total_potential_cells += potential_cells
        total_tagged_cells += table_analysis['cells_with_tags']
        analysis['summary']['total_tags'] += table_analysis['tag_count']
    
    # Общее покрытие
    if total_potential_cells > 0:
        analysis['summary']['coverage_percent'] = round(
            (total_tagged_cells / total_potential_cells) * 100, 2
        )
    
    return analysis


def print_analysis(analysis: Dict):
    """Красиво выводит результаты анализа."""
    
    print("\n" + "="*80)
    print("📊 АНАЛИЗ ЗАПОЛНЕНИЯ ТЕГОВ В ДОКУМЕНТЕ")
    print("="*80)
    
    summary = analysis['summary']
    print(f"\n📈 ОБЩАЯ СТАТИСТИКА:")
    print(f"   Всего таблиц:              {analysis['total_tables']}")
    print(f"   С тегами:                  {summary['tables_with_tags']}")
    print(f"   Без тегов:                 {summary['tables_without_tags']}")
    print(f"   Полностью заполнены:       {summary['tables_fully_tagged']}")
    print(f"   Частично заполнены:        {summary['tables_partially_tagged']}")
    print(f"   Всего тегов:               {summary['total_tags']}")
    print(f"   Общее покрытие:            {summary['coverage_percent']}%")
    
    # Таблицы без тегов
    tables_without_tags = [t for t in analysis['tables'] if t['status'] == 'NOT_TAGGED']
    if tables_without_tags:
        print(f"\n❌ ТАБЛИЦЫ БЕЗ ТЕГОВ ({len(tables_without_tags)}):")
        for t in tables_without_tags:
            print(f"   Таблица {t['index']:2d}: {t['rows']} строк × {t['cols']} столбцов | "
                  f"Всего ячеек: {t['potential_cells']}")
            if t['issues']:
                print(f"                   ⚠️ {', '.join(t['issues'])}")
    
    # Таблицы с частичным покрытием
    tables_partial = [t for t in analysis['tables'] if t['status'] == 'PARTIALLY_TAGGED']
    if tables_partial:
        print(f"\n⚠️  ТАБЛИЦЫ С ЧАСТИЧНЫМ ПОКРЫТИЕМ ({len(tables_partial)}):")
        for t in tables_partial:
            print(f"   Таблица {t['index']:2d}: {t['coverage_percent']:5.1f}% ({t['cells_with_tags']}/{t['potential_cells']} ячеек) | "
                  f"{t['tag_count']} тегов")
            if t['tags_sample']:
                print(f"                   Примеры: {', '.join(t['tags_sample'][:3])}")
            if t['issues']:
                print(f"                   ⚠️ {', '.join(t['issues'])}")
    
    # Таблицы с полным покрытием
    tables_full = [t for t in analysis['tables'] if t['status'] == 'FULLY_TAGGED']
    if tables_full:
        print(f"\n✅ ТАБЛИЦЫ С ПОЛНЫМ ПОКРЫТИЕМ ({len(tables_full)}):")
        for t in tables_full[:10]:  # Показываем первые 10
            print(f"   Таблица {t['index']:2d}: 100% ({t['cells_with_tags']}/{t['potential_cells']} ячеек) | "
                  f"{t['tag_count']} тегов")
        if len(tables_full) > 10:
            print(f"   ... и еще {len(tables_full) - 10} таблиц с полным покрытием")
    
    # Детальная информация по всем таблицам
    print(f"\n📋 ДЕТАЛЬНАЯ ИНФОРМАЦИЯ ПО ВСЕМ ТАБЛИЦАМ:")
    print("-" * 120)
    print(f"{'№':<3} {'Статус':<15} {'Строк':<6} {'Столб':<6} {'Ячеек':<8} "
          f"{'Покрыто':<10} {'Теги':<6} {'Пусто':<6} {'Прочерк':<8} {'Проблемы':<30}")
    print("-" * 120)
    
    for t in analysis['tables']:
        status = t['status']
        status_icon = {'FULLY_TAGGED': '✅', 'PARTIALLY_TAGGED': '⚠️ ', 'NOT_TAGGED': '❌'}[status]
        
        issues_str = ', '.join(t['issues'][:2]) if t['issues'] else '-'
        if len(t['issues']) > 2:
            issues_str += f" (+{len(t['issues']) - 2})"
        
        print(f"{t['index']:<3} {status_icon} {status:<12} {t['rows']:<6} {t['cols']:<6} "
              f"{t['potential_cells']:<8} {t['coverage_percent']:<9}% {t['tag_count']:<6} "
              f"{t['empty_cells']:<6} {t['cells_with_dash']:<8} {issues_str[:28]:<30}")
    
    print("-" * 120)
    
    print("\n" + "="*80)


def generate_detailed_report(analysis: Dict, output_path: str):
    """Генерирует детальный текстовый отчет."""
    
    lines = []
    lines.append("="*80)
    lines.append("АНАЛИЗ ЗАПОЛНЕНИЯ ТЕГОВ В DOCX - ДЕТАЛЬНЫЙ ОТЧЕТ")
    lines.append("="*80)
    lines.append("")
    
    summary = analysis['summary']
    lines.append("ОБЩАЯ СТАТИСТИКА")
    lines.append("-" * 40)
    lines.append(f"Всего таблиц в документе:          {analysis['total_tables']}")
    lines.append(f"Таблицы с тегами:                  {summary['tables_with_tags']}")
    lines.append(f"Таблицы без тегов:                 {summary['tables_without_tags']}")
    lines.append(f"Таблицы с полным покрытием:        {summary['tables_fully_tagged']}")
    lines.append(f"Таблицы с частичным покрытием:     {summary['tables_partially_tagged']}")
    lines.append(f"Всего вставлено тегов:             {summary['total_tags']}")
    lines.append(f"Среднее покрытие:                  {summary['coverage_percent']}%")
    lines.append("")
    
    # Список таблиц без тегов
    tables_without_tags = [t for t in analysis['tables'] if not t['has_tags']]
    if tables_without_tags:
        lines.append("ТАБЛИЦЫ БЕЗ ТЕГОВ")
        lines.append("-" * 40)
        for t in tables_without_tags:
            lines.append(f"Таблица {t['index']}")
            lines.append(f"  Размер: {t['rows']} строк × {t['cols']} столбцов")
            lines.append(f"  Ячеек: {t['potential_cells']}")
            lines.append(f"  Пусто: {t['empty_cells']}")
            lines.append(f"  Содержит текст: {t['cells_with_text']}")
            lines.append(f"  Содержит прочерки: {t['cells_with_dash']}")
            if t['issues']:
                lines.append(f"  Проблемы: {'; '.join(t['issues'])}")
            lines.append("")
    
    # Таблицы с частичным покрытием
    tables_partial = [t for t in analysis['tables'] if t['coverage_percent'] > 0 and t['coverage_percent'] < 100]
    if tables_partial:
        lines.append("ТАБЛИЦЫ С ЧАСТИЧНЫМ ПОКРЫТИЕМ")
        lines.append("-" * 40)
        for t in tables_partial:
            lines.append(f"Таблица {t['index']}")
            lines.append(f"  Размер: {t['rows']} строк × {t['cols']} столбцов")
            lines.append(f"  Покрытие: {t['coverage_percent']}% ({t['cells_with_tags']}/{t['potential_cells']} ячеек)")
            lines.append(f"  Тегов: {t['tag_count']}")
            lines.append(f"  Примеры тегов: {', '.join(t['tags_sample'][:5])}")
            lines.append(f"  Пусто: {t['empty_cells']}")
            lines.append(f"  Прочерки: {t['cells_with_dash']}")
            if t['issues']:
                lines.append(f"  Проблемы: {'; '.join(t['issues'])}")
            lines.append("")
    
    # Таблицы с полным покрытием
    tables_full = [t for t in analysis['tables'] if t['coverage_percent'] == 100]
    if tables_full:
        lines.append("ТАБЛИЦЫ С ПОЛНЫМ ПОКРЫТИЕМ")
        lines.append("-" * 40)
        lines.append(f"Всего: {len(tables_full)} таблиц")
        lines.append("Номера таблиц: " + ", ".join(str(t['index']) for t in tables_full))
        lines.append("")
    
    # Выводы и рекомендации
    lines.append("ВЫВОДЫ И РЕКОМЕНДАЦИИ")
    lines.append("-" * 40)
    
    if summary['tables_without_tags'] > 0:
        lines.append(f"⚠️  Обнаружено {summary['tables_without_tags']} таблиц без тегов.")
        lines.append("    РЕКОМЕНДАЦИЯ: Проверьте эти таблицы в шаблоне - они не будут заполняться.")
    
    if summary['tables_partially_tagged'] > 0:
        lines.append(f"⚠️  Обнаружено {summary['tables_partially_tagged']} таблиц с частичным покрытием.")
        lines.append("    РЕКОМЕНДАЦИЯ: Добавьте теги в оставшиеся ячейки или убедитесь, что это намеренно.")
    
    if summary['coverage_percent'] < 50:
        lines.append("🔴 НИЗКОЕ ОБЩЕЕ ПОКРЫТИЕ - менее 50%!")
        lines.append("   Большинство таблиц не подготовлены к автоматическому заполнению.")
    elif summary['coverage_percent'] < 80:
        lines.append("🟡 СРЕДНЕЕ ПОКРЫТИЕ - от 50% до 80%.")
        lines.append("   Некоторые таблицы требуют внимания.")
    else:
        lines.append("🟢 ХОРОШЕЕ ПОКРЫТИЕ - более 80%.")
        lines.append("   Большинство таблиц готовы к заполнению.")
    
    lines.append("")
    
    report_text = "\n".join(lines)
    Path(output_path).parent.mkdir(parents=True, exist_ok=True)
    with open(output_path, 'w', encoding='utf-8') as f:
        f.write(report_text)
    
    return report_text


def main():
    """Главная функция."""
    
    # Путь к документу
    doc_path = "input/Бюллетень_ШАБЛОН_С_ТЕГАМИ_v2.docx"
    
    if not Path(doc_path).exists():
        print(f"❌ Файл не найден: {doc_path}")
        return
    
    print(f"📄 Анализируем: {doc_path}")
    
    # Анализируем
    analysis = analyze_document_structure(doc_path)
    
    # Выводим результаты
    print_analysis(analysis)
    
    # Сохраняем отчет
    report_path = "output/analysis_tag_coverage.txt"
    report_text = generate_detailed_report(analysis, report_path)
    print(f"\n✅ Отчет сохранен: {report_path}")
    
    return analysis


if __name__ == "__main__":
    main()
