# confic.py
import csv
import re
from collections import Counter

import pandas as pd
from docx.oxml.ns import qn
from docx.table import _Cell, Table

# === Константы ===
EMPTY_CELL_MARKER = "—"
OKVED_CODE_COLUMN_INDEX = 0
OKVED_NAME_COLUMN_INDEX = 1


# === Утилиты ===
def _normalize_text(text: str) -> str:
    """
    Базовая нормализация: очистка пробелов и приведение к нижнему регистру.
    """
    if not text:
        return ""
    # Приводим к нижнему регистру и нормализуем пробелы
    text = text.lower().strip()
    # Заменяем ё на е (для русского языка)
    text = text.replace('ё', 'е')
    # Удаляем множественные пробелы
    text = re.sub(r'\s+', ' ', text).strip()
    return text


def canonical_okved(code: str) -> str:
    """
    Приводит код ОКВЭД к единому каноническому виду.
    Правила:
    1. Убираем пробелы по краям.
    2. Приводим к верхнему регистру.
    Никакой магии - только очистка от человеческого фактора.
    """
    if not code:
        return ""
    return str(code).strip().upper()


def get_cleaned_cell_text(cell: _Cell) -> str:
    return ' '.join(p.text.replace('\n', ' ').strip() for p in cell.paragraphs).strip()


def looks_like_data_value(text: str) -> bool:
    """Похоже ли содержимое ячейки на числовое ЗНАЧЕНИЕ, а не на подпись.

    Очистка ячеек раньше ориентировалась на «есть хоть одна цифра» и из-за
    этого затирала прочерком заголовки, в которых цифра тоже есть:
    «Продолжение таблицы 1.», «Коэффициент текущей ликвидности до < 100%»,
    «длительность 1 оборота». В готовом бюллетене на их месте оставался
    прочерк, то есть терялся текст, который вообще не имеет отношения к
    данным.

    Значение от подписи надёжно отличает отсутствие букв: во всём исходном
    бюллетене нет ни одной ячейки-данных с буквами (проверено по всем 140
    таблицам — все 30 таких текстов оказались подписями).
    """
    return any(ch.isdigit() for ch in text) and not any(ch.isalpha() for ch in text)


def dominant_row_bold(row, exclude_cell=None) -> "Optional[bool]":
    """Определяет преобладающую жирность среди непустых ячеек строки данных.

    Строки бюллетеня обычно жирные целиком (числа + прочерки одинаково),
    но встречаются реальные дефекты исходного документа: одна ячейка в
    строке набрана без жирности, хотя остальные — жирные (см.
    НАЙДЕННЫЕ_ДЕФЕКТЫ_ШАБЛОНА.md — тот же класс несогласованного
    оформления, что там уже нашли для "-"). Большинство соседей по строке —
    надёжный сигнал, каким должно быть форматирование той единственной
    ячейки, что выбивается.

    exclude_cell: ячейку саму себя (обычно ту, что сейчас заполняем)
    исключаем из подсчёта — иначе её собственная (возможно, как раз
    дефектная) жирность перевесит счёт себя же.
    """
    votes = []
    for cell in row.cells:
        if exclude_cell is not None and cell._tc is exclude_cell._tc:
            continue
        for p in cell.paragraphs:
            for run in p.runs:
                if run.text.strip():
                    votes.append(bool(run.bold))
                    break
            else:
                continue
            break
    if not votes:
        return None
    return sum(votes) > len(votes) / 2


def set_paragraph_text_keep_format(paragraph, text: str) -> None:
    """Заменяет текст параграфа, СОХРАНЯЯ форматирование (шрифт, размер,
    жирность и т.д.) уже существующего в нём текста.

    Обычное `paragraph.text = text` (см. python-docx) внутри делает
    `self.clear()` + `self.add_run(text)` — то есть удаляет ВСЕ существующие
    runs вместе с их форматированием и создаёт новый run с форматированием
    по умолчанию (обычно это даёт другой шрифт/размер, чем у остальной
    таблицы: например, Calibri 11 вместо Times New Roman 12, как в
    исходном бюллетене). Из-за этого после вставки тегов и заполнения их
    данными числа визуально отличались от остального документа.

    Вместо этого пишем текст в ПЕРВЫЙ существующий run (его форматирование
    остаётся как было), а все прочие runs этого параграфа опустошаем — так
    старый текст не дублируется, а форматирование не теряется. Если runs в
    параграфе вообще нет (пустой параграф), тогда действительно приходится
    создать новый — сохранять там нечего.
    """
    runs = paragraph.runs
    if runs:
        runs[0].text = text
        for extra_run in runs[1:]:
            extra_run.text = ""
    else:
        paragraph.add_run(text)


def normalize_dash_bold_in_document(doc) -> int:
    """Проходит по ВСЕМ таблицам готового документа и выравнивает жирность
    ячеек-прочерков ("-") по преобладающей жирности остальной строки.

    Точечные правки при заполнении тегов (see fill_word_template_by_tags_v2,
    generate_word_template) чинят только те прочерки, которые сами
    записываем в процессе работы программы. Но встречаются прочерки,
    которые были жирностью None ещё в исходном input/Бюллетень.docx и
    которые программа вообще не трогает (в них никогда не было тега —
    это не заполняемая ячейка, а изначально статичный "-"). Такой прочерк
    иначе так и остаётся несогласованным с соседями по строке в готовом
    документе. Этот проход ловит оба случая разом, независимо от того,
    когда и кем был записан текст ячейки.
    """
    fixed = 0
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                if get_cleaned_cell_text(cell) != "-":
                    continue
                dominant = dominant_row_bold(row, exclude_cell=cell)
                if dominant is None:
                    continue
                for p in cell.paragraphs:
                    for run in p.runs:
                        if run.text.strip() and bool(run.bold) != dominant:
                            run.bold = dominant
                            fixed += 1
    return fixed


def _table_declared_width_twips(table):
    """Заявленная ширина таблицы в twips — по сумме столбцов сетки
    (w:tblGrid), а если сетки нет, то по w:tblW (только если задан в dxa,
    то есть в абсолютных единицах — проценты тут ни при чём)."""
    tbl = table._tbl
    grid = tbl.find(qn('w:tblGrid'))
    if grid is not None:
        total = 0
        found = False
        for gridCol in grid.findall(qn('w:gridCol')):
            w = gridCol.get(qn('w:w'))
            if w:
                total += int(w)
                found = True
        if found:
            return total
    tblW = tbl.tblPr.find(qn('w:tblW'))
    if tblW is not None and tblW.get(qn('w:type')) == 'dxa':
        w = tblW.get(qn('w:w'))
        if w:
            return int(w)
    return None


def autofit_tables_to_window(doc) -> int:
    """Включает режим Word "Автоподбор по ширине окна" (Таблица → Свойства
    таблицы → Автоподбор → Автоподбор по ширине окна) ТОЛЬКО для тех
    таблиц, которые реально шире печатной области страницы.

    В исходном бюллетене часть таблиц задана фиксированной шириной (dxa),
    которая местами шире печатной области (обнаружено: сетка таблицы
    ~14790 twips против ~14570 twips полезной ширины страницы — таблица
    выходит за правый край примерно на 0.15 дюйма). Именно для таких
    переключаем ширину на 100% окна (w:tblW type=pct) и включаем
    автоподбор (w:tblLayout type=autofit) — пропорции столбцов друг
    относительно друга сохраняются, вся таблица масштабируется под
    печатную область.

    Таблицы, которые и так укладываются в ширину страницы (например, уже
    вручную подогнанные), НЕ трогаем — иначе растянули бы их на все окно
    вместо того, чтобы просто убрать имеющееся переполнение.
    """
    sec = doc.sections[0]
    usable_twips = (sec.page_width - sec.left_margin - sec.right_margin - sec.gutter) / 635

    changed = 0
    for table in doc.tables:
        declared_twips = _table_declared_width_twips(table)
        if declared_twips is None or declared_twips <= usable_twips:
            continue

        tblPr = table._tbl.tblPr
        tblW = tblPr.find(qn('w:tblW'))
        if tblW is None:
            tblW = tblPr.makeelement(qn('w:tblW'), {})
            jc = tblPr.find(qn('w:jc'))
            if jc is not None:
                jc.addprevious(tblW)
            else:
                tblPr.append(tblW)
        tblW.set(qn('w:type'), 'pct')
        tblW.set(qn('w:w'), '5000')
        table.autofit = True
        changed += 1
    return changed


_ROBUST_ENCODINGS = ['utf-8-sig', 'windows-1251', 'cp1251', 'utf-8', 'latin1']


def _read_csv_rows_robustly(filepath, delimiter=';', quotechar='"'):
    """Читает CSV-файл построчно (список списков-ячеек), перебирая
    кодировки — на случай, если файл пересохранили из Excel на Windows
    (обычно даёт windows-1251, а не UTF-8). Аналог _read_csv_robustly, но
    без pandas — для мест, где нужен именно "сырой" csv.reader.
    """
    last_error = None
    for encoding in _ROBUST_ENCODINGS:
        try:
            with open(filepath, encoding=encoding, newline='') as f:
                reader = csv.reader(f, delimiter=delimiter, quotechar=quotechar)
                return [row for row in reader]
        except UnicodeDecodeError as e:
            last_error = e
            continue
    raise ValueError(f"❌ Не удалось прочитать {filepath} ни с одной кодировкой ({last_error}).")


def _read_csv_robustly(filepath, header_row=0):
    encodings = ['utf-8-sig', 'windows-1251', 'cp1251', 'utf-8', 'latin1']
    for encoding in encodings:
        try:
            print(f"Попытка чтения {filepath} с кодировкой {encoding}")
            df = pd.read_csv(filepath, sep=';', encoding=encoding, header=header_row)
            print(f"✅ Успешно прочитано с кодировкой {encoding}")
            if df.shape[1] == 1:
                print(
                    f"⚠️ ВНИМАНИЕ: CSV-файл прочитан как один столбец. Возможно, проблема с разделителями или кавычками.")
                print(f"📋 Заголовок: {df.columns[0]}")
                print("💡 Проверь, что файл сохранён с разделителями ';' и без двойных кавычек вокруг всей строки.")
            return df
        except UnicodeDecodeError as e:
            print(f"❌ Ошибка кодировки {encoding}: {e}")
        except Exception as e:
            print(f"❌ Другая ошибка с {encoding}: {e}")
    raise ValueError("❌ Не удалось прочитать CSV-файл ни с одной кодировкой.")


# === Загрузка справочников ===
def load_okved_map(filepath):
    print(f"Загрузка справочника ОКВЭД из: {filepath}")
    df = _read_csv_robustly(filepath, header_row=None)
    # .iloc — это позиционный индексатор pandas.→ взять все строки.→ взять только первые два столбца (с индексами 0
    # и 1, правая граница не включается).
    # Полученная таблица перезаписывает df. Лишние колонки отбрасываются.
    df = df.iloc[:, :2]
    # -----

    df.columns = ['Код', 'Наименование']
    df = df.dropna(subset=['Код'])
    df['Код'] = df['Код'].astype(str).str.strip()
    df['Наименование'] = df['Наименование'].astype(str).str.strip()
    df['Наименование_норм'] = df['Наименование'].apply(_normalize_text)
    okved_to_name = dict(zip(df['Код'], df['Наименование']))
    name_to_okved_cleaned = dict(zip(df['Наименование_норм'], df['Код']))
    print(f"📘 Загружено {len(df)} записей ОКВЭД.")
    return okved_to_name, name_to_okved_cleaned


def load_table_source_map(filepath):
    print(f"Загрузка сопоставления таблиц из: {filepath}")
    mapping_dict = {}

    rows = [row for row in _read_csv_rows_robustly(filepath, delimiter=';') if row]

    if not rows:
        print("⚠️ Файл маппинга пуст.")
        return mapping_dict

    header = [str(col).strip().lower() for col in rows[0]]
    if 'таблица' in header and ('файл' in header or 'источник' in header):
        data_rows = rows[1:]
    else:
        # Если заголовки не обнаружены, возможно файл читается без заголовка
        data_rows = rows

    for row in data_rows:
        if len(row) < 2:
            continue
        raw_name = str(row[0]).strip().strip('"').strip()
        raw_src = str(row[1]).strip().strip('"').strip()
        if not raw_name or not raw_src:
            continue
        norm_key = _normalize_text(raw_name)
        mapping_dict[norm_key] = raw_src

    print(f"📘 Загружено сопоставлений: {len(mapping_dict)}")
    return mapping_dict


def load_column_mapping(filepath):
    """
    Загружает column_mapping.csv и возвращает:
    - word_to_indicator: название показателя → код индикатора
    - indicator_to_excel: код индикатора → (номер колонки 2022, номер колонки 2023)
    - indicator_to_file: код индикатора → имя Excel-файла
    - file_word_to_indicator: (файл, нормализованное название) → код индикатора
    """
    word_to_indicator = {}
    indicator_to_excel = {}
    indicator_to_file = {}
    file_word_to_indicator = {}

    with open(filepath, encoding="utf-8-sig") as f:
        reader = csv.reader(f, delimiter=';')
        headers = next(reader)  # Пропускаем заголовок

        duplicate_indicator_names = {}

        for row in reader:
            if len(row) < 5:
                continue  # Пропускаем неполные строки

            excel_file = str(row[0]).strip()
            word_name = str(row[1]).strip()
            indicator = str(row[2]).strip()
            col_22 = str(row[3]).strip() if len(row) > 3 else ''
            col_23 = str(row[4]).strip() if len(row) > 4 else ''

            if col_22.lower() == 'nan':
                col_22 = ''
            if col_23.lower() == 'nan':
                col_23 = ''

            word_to_indicator[word_name] = indicator
            indicator_to_excel[indicator] = (col_22, col_23)
            indicator_to_file[indicator] = excel_file
            # NEW: Added mapping per file to avoid conflicts when same indicator in multiple files
            file_word_to_indicator[(excel_file, _normalize_text(word_name))] = indicator

            duplicate_indicator_names.setdefault(indicator, set()).add(word_name)

    for indicator, names in duplicate_indicator_names.items():
        if len(names) > 1:
            print(f"⚠️ Дублирующийся код индикатора '{indicator}' для разных названий: {sorted(names)}")

    return word_to_indicator, indicator_to_excel, indicator_to_file, file_word_to_indicator


# === Поиск названия таблицы ===
def get_table_name(table: Table, known_table_names, return_details: bool = False):
    """Находит название таблицы по параграфам выше неё.

    Поддерживает:
    - Точные совпадения (приоритет)
    - Частичные совпадения (одно в другом)
    - Нечеткие совпадения через fuzzy match (для опечаток)

    return_details=True возвращает (title, saw_heading_text) вместо
    просто title: saw_heading_text говорит, была ли НАД таблицей хоть
    какая-то содержательная (не пустая, не техническая) строка-заголовок —
    даже если она ни с чем не совпала. Это отличает две разные ситуации:
    "заголовка вообще нет" (похоже на продолжение предыдущей таблицы без
    явной метки) от "заголовок есть, но это явно ДРУГАЯ, незнакомая
    таблица" — вызывающий код должен обработать их по-разному.
    """
    from difflib import SequenceMatcher

    ignore_phrases = ['продолжение таблицы', 'таблица', 'график', 'рис.', 'рис', 'по всей форме', 'приложение', 'форма',
                      'лист', 'страница', 'тысяч рублей', 'на конец года']
    # Обычная проверка "phrase in text" ловит фразу как ПОДСТРОКУ где угодно —
    # из-за этого одиночное слово "форма" ложно срабатывает на настоящих
    # заголовках вроде "...ПО ОРГАНИЗАЦИОННО-ПРАВОВЫМ ФОРМАМ" или "...ПО ФОРМАМ
    # СОБСТВЕННОСТИ" (там "форма" — просто начало слова "формам", а не сама
    # фраза). Границы слова (\b) отсекают такие случаи, оставляя прежнее
    # поведение для реальных технических пометок ("Форма 1", "рис. 2" и т.п.).
    ignore_phrase_patterns = [re.compile(r'\b' + re.escape(phrase) + r'\b') for phrase in ignore_phrases]

    def _has_ignore_phrase(text_norm: str) -> bool:
        return any(pattern.search(text_norm) for pattern in ignore_phrase_patterns)

    # Собираем все параграфы выше таблицы, чтобы выбрать наиболее подходящий заголовок.
    exact_matches = []  # Точные совпадения (имеют приоритет)
    partial_matches = []  # Частичные совпадения
    fuzzy_matches = []  # Нечеткие совпадения (для опечаток)
    saw_heading_text = False
    para_count = 0
    MAX_PARAS_TO_CHECK = 20
    prev_elem = table._element.getprevious()
    
    # Накопимаем параграфы для объединения многострочных заголовков
    para_buffer = []

    while prev_elem is not None and para_count < MAX_PARAS_TO_CHECK:
        if prev_elem.tag.endswith('p'):
            text = (prev_elem.text or "").strip()
            prev_elem = prev_elem.getprevious()
            para_count += 1
            
            if not text:
                # Если встретили пустой параграф, проверяем накопленный буфер
                if para_buffer:
                    combined_text = " ".join(reversed(para_buffer))
                    combined_norm = _normalize_text(combined_text)
                    if not _has_ignore_phrase(combined_norm):
                        for known_title in known_table_names:
                            known_norm = _normalize_text(known_title)
                            if combined_norm == known_norm:
                                score = len(known_norm) / (para_count + 1)
                                exact_matches.append((score, known_title, para_count))
                para_buffer = []
                continue

            text_norm = _normalize_text(text)
            if _has_ignore_phrase(text_norm):
                para_buffer = []
                continue

            saw_heading_text = True
            # Добавляем параграф в буфер
            para_buffer.append(text)
            
            # Проверяем каждый параграф отдельно
            for known_title in known_table_names:
                known_norm = _normalize_text(known_title)
                if text_norm == known_norm:
                    # Точное совпадение
                    score = len(known_norm) / (para_count + 1)
                    exact_matches.append((score, known_title, para_count))
                elif text_norm in known_norm or known_norm in text_norm:
                    # Частичное совпадение
                    score = len(known_norm) / (para_count + 1)
                    partial_matches.append((score, known_title, para_count))
                else:
                    # Нечеткое совпадение (для опечаток)
                    similarity = SequenceMatcher(None, text_norm, known_norm).ratio()
                    if similarity >= 0.80:  # Порог для нечеткого совпадения
                        fuzzy_matches.append((similarity * len(known_norm) / (para_count + 1), known_title, para_count))
            
            # Проверяем объединенные параграфы (буфер из последних 2 параграфов)
            if len(para_buffer) >= 2:
                combined_text = " ".join(reversed(para_buffer[:2]))
                combined_norm = _normalize_text(combined_text)
                for known_title in known_table_names:
                    known_norm = _normalize_text(known_title)
                    if combined_norm == known_norm:
                        # Точное совпадение после объединения
                        score = len(known_norm) / (para_count + 1)
                        exact_matches.append((score, known_title, para_count))

        else:
            prev_elem = prev_elem.getprevious()
            para_count += 1
            para_buffer = []  # Сбрасываем буфер при встречке не-параграфа

    def _result(title):
        return (title, saw_heading_text) if return_details else title

    # Приоритет: точные совпадения > частичные совпадения > нечеткие совпадения
    if exact_matches:
        exact_matches.sort(reverse=True)
        best_title = exact_matches[0][1]
        print(f"🔍 Заголовок найден (точное совпадение): {best_title}")
        return _result(best_title)

    if partial_matches:
        partial_matches.sort(reverse=True)
        best_title = partial_matches[0][1]
        print(f"🔍 Заголовок найден (частичное совпадение): {best_title}")
        return _result(best_title)

    if fuzzy_matches:
        fuzzy_matches.sort(reverse=True)
        best_title = fuzzy_matches[0][1]
        print(f"🔍 Заголовок найден (нечеткое совпадение): {best_title}")
        return _result(best_title)

    print("⚠️ Заголовок не распознан")
    return _result(None)


# === Загрузка Excel-данных ===
def get_excel_data(excel_path, okved_codes_set):
    try:
        df = pd.read_excel(excel_path, header=None, dtype=str)
    except Exception as e:
        print(f"❌ Ошибка чтения Excel: {e}")
        return {}

    header_map = {}
    header_row_index = -1

    # 🔍 Поиск строки заголовков (обычно строка с 'А' и '1')
    for i, row in df.head(15).iterrows():
        if len(row) > 2 and pd.notna(row.iloc[0]) and pd.notna(row.iloc[2]):
            if str(row.iloc[0]).strip().upper() == 'А' and str(row.iloc[2]).strip() == '1':
                header_row_index = i
                break

    if header_row_index == -1:
        print(f"⚠️ Не найдена строка заголовков в {excel_path.name}")
        return {}

    header_row = df.iloc[header_row_index]
    for i, val in enumerate(header_row):
        if pd.notna(val) and i >= 2:
            header_map[i] = str(val).strip()

    data_start_row = header_row_index + 1
    data_dict = {}

    for _, row in df.iloc[data_start_row:].iterrows():
        if row.empty or pd.isna(row.iloc[OKVED_CODE_COLUMN_INDEX]):
            continue

        okved_code_raw = str(row.iloc[OKVED_CODE_COLUMN_INDEX]).strip()
        if not okved_code_raw:
            continue
        
        # Применяем канонизацию кода ОКВЭД
        okved_code = canonical_okved(okved_code_raw)

        # 🔍 Если фильтрация включена — пропускаем лишние
        if okved_codes_set and okved_code not in okved_codes_set:
            continue

        data_dict[okved_code] = {}

        for col_idx, key_from_header in header_map.items():
            if col_idx >= len(row):
                continue
            value = row.iloc[col_idx]
            if pd.isna(value) or str(value).strip() in ('-', '""', ''):
                formatted_value = EMPTY_CELL_MARKER
            else:
                try:
                    num_value = float(str(value).replace(',', '.').replace(' ', ''))
                    if num_value == int(num_value):
                        formatted_value = f"{int(num_value):,}".replace(',', ' ')
                    else:
                        formatted_value = f"{num_value:,.2f}".replace(',', ' ').replace('.', ',')
                except (ValueError, TypeError):
                    formatted_value = str(value).strip()
            data_dict[okved_code][str(col_idx)] = formatted_value

    if not data_dict:
        print(f"⚠️ В Excel-файле {excel_path.name} не найдено ни одного подходящего кода ОКВЭД.")
    else:
        print(f"📊 Загружено данных для {len(data_dict)} кодов ОКВЭД из {excel_path.name}")

    return data_dict


# Короткие предлоги/союзы, чьё отсутствие или добавление в тексте ячейки
# (например, потерянное "и" при переносе строки в исходном docx) не должно
# мешать сопоставлению с эталонным названием из справочника.
_SINGLE_WORD_DIFF_ALLOWED = {
    'и', 'а', 'но', 'или', 'да', 'не', 'ни', 'то', 'же', 'ли',
    'с', 'со', 'в', 'во', 'к', 'ко', 'у', 'о', 'об', 'обо',
    'от', 'до', 'из', 'изо', 'за', 'на', 'по', 'под', 'над',
    'при', 'для', 'без', 'через',
}


def _match_by_single_preposition_diff(text_norm, name):
    """
    Возвращает True, если text_norm и name отличаются ровно одним словом,
    и это слово — короткий предлог/союз (см. _SINGLE_WORD_DIFF_ALLOWED).
    """
    text_words = Counter(text_norm.split())
    name_words = Counter(name.split())
    diff = (text_words - name_words) + (name_words - text_words)
    extra_words = list(diff.elements())
    if len(extra_words) != 1:
        return False
    return extra_words[0] in _SINGLE_WORD_DIFF_ALLOWED


# === Поиск кода ОКВЭД по названию ===
def find_okved_code(cell_text, name_to_okved_cleaned):
    text_norm = _normalize_text(cell_text)
    if not text_norm or text_norm in ('код', 'наименование'):
        return None
    if text_norm in name_to_okved_cleaned:
        return name_to_okved_cleaned[text_norm]
    for name, code in name_to_okved_cleaned.items():
        if len(text_norm) >= 3 and (text_norm in name or name in text_norm):
            print(f"⚠️ Частичное совпадение: '{cell_text}' ~ '{name}' → {code}")
            return code
    for name, code in name_to_okved_cleaned.items():
        if _match_by_single_preposition_diff(text_norm, name):
            print(f"⚠️ Совпадение с точностью до предлога/союза: '{cell_text}' ~ '{name}' → {code}")
            return code
    return None
