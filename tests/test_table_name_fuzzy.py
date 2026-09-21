import tempfile
import unittest
from pathlib import Path

from config import get_table_name
from docx import Document


class TestTableNameFuzzy(unittest.TestCase):
    def test_get_table_name_fuzzy_from_paragraphs(self):
        with tempfile.TemporaryDirectory() as tmpdir:
            doc_path = Path(tmpdir) / "test.docx"
            doc = Document()
            doc.add_paragraph("Раздел 1")
            doc.add_paragraph("1. Баланс организаций по видам экономической деятельности")
            doc.add_table(rows=1, cols=1)
            doc.save(doc_path)

            loaded = Document(doc_path)
            table = loaded.tables[0]
            known_titles = [
                "1. Баланс организаций по видам экономической деятельности",
                "2. Отчет о движении денежных средств"
            ]

            title = get_table_name(table, known_titles)
            self.assertEqual(title, known_titles[0])

    def test_get_table_name_fuzzy_with_typo(self):
        with tempfile.TemporaryDirectory() as tmpdir:
            doc_path = Path(tmpdir) / "test_typo.docx"
            doc = Document()
            doc.add_paragraph("1. Баланс оргнизаций по видам экономики")
            doc.add_table(rows=1, cols=1)
            doc.save(doc_path)

            loaded = Document(doc_path)
            table = loaded.tables[0]
            known_titles = [
                "1. Баланс организаций по видам экономической деятельности"
            ]

            title = get_table_name(table, known_titles)
            self.assertEqual(title, known_titles[0])

    def test_return_details_flags_unrecognized_heading_as_seen(self):
        # Заголовок ЕСТЬ (это не "продолжение таблицы без метки"), но он не
        # совпадает ни с одним известным названием — похоже на совсем
        # другую, незнакомую таблицу. logic.py использует этот флаг, чтобы
        # не наследовать источник данных предыдущей таблицы вслепую.
        with tempfile.TemporaryDirectory() as tmpdir:
            doc_path = Path(tmpdir) / "test_foreign.docx"
            doc = Document()
            doc.add_paragraph("99. Средняя температура воздуха по регионам")
            doc.add_table(rows=1, cols=1)
            doc.save(doc_path)

            loaded = Document(doc_path)
            table = loaded.tables[0]
            known_titles = ["1. Баланс организаций по видам экономической деятельности"]

            title, saw_heading_text = get_table_name(table, known_titles, return_details=True)
            self.assertIsNone(title)
            self.assertTrue(saw_heading_text)

    def test_return_details_no_heading_at_all(self):
        # Совсем без заголовка над таблицей (типичный случай настоящего
        # "молчаливого" продолжения предыдущей таблицы) — saw_heading_text
        # должен быть False, чтобы logic.py по-прежнему мог унаследовать
        # источник предыдущей таблицы.
        with tempfile.TemporaryDirectory() as tmpdir:
            doc_path = Path(tmpdir) / "test_no_heading.docx"
            doc = Document()
            doc.add_table(rows=1, cols=1)
            doc.save(doc_path)

            loaded = Document(doc_path)
            table = loaded.tables[0]
            known_titles = ["1. Баланс организаций по видам экономической деятельности"]

            title, saw_heading_text = get_table_name(table, known_titles, return_details=True)
            self.assertIsNone(title)
            self.assertFalse(saw_heading_text)


if __name__ == "__main__":
    unittest.main()
