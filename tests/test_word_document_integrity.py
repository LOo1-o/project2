"""Что программа имеет право менять в документе, а что обязана сохранить.

Все проверки здесь выросли из дефектов, найденных сверкой готового
бюллетеня с заполненным вручную оригиналом по всем 140 таблицам.
"""
import unittest

from docx import Document
from docx.shared import Twips

from config import (
    autofit_tables_to_window,
    dominant_row_bold,
    looks_like_data_value,
    normalize_dash_bold_in_document,
)
from data_filler_v2 import fill_word_template_by_tags_v2


class TestDataValueDetection(unittest.TestCase):
    """Очистка ячеек раньше ориентировалась на «есть хоть одна цифра» и
    затирала прочерком подписи, в которых цифра тоже есть. Тега для них
    нет, поэтому текст уже не возвращался — в готовом документе на его
    месте оставался прочерк."""

    def test_headings_with_digits_are_not_data(self):
        for text in [
            "Продолжение таблицы 1.",
            "Продолжение таблицы 20.",
            "Коэффициент текущей ликвидности до < 100%",
            "Коэффициент текущей ликвидности от > = 100% до < 200%",
            "длительность  1 оборота",
            "Соотношение заемных и собственных средств (нормальное ограничение 100%)",
        ]:
            with self.subTest(text=text):
                self.assertFalse(looks_like_data_value(text))

    def test_numbers_are_data(self):
        for text in ["94", "40,7", "-675 642", "1 637 622", "86 180,6", "-2,0"]:
            with self.subTest(text=text):
                self.assertTrue(looks_like_data_value(text))

    def test_text_without_digits_is_not_data(self):
        for text in ["", "-", "Всего по обследуемым видам", "х"]:
            with self.subTest(text=text):
                self.assertFalse(looks_like_data_value(text))


class TestMissingYearIsNotSubstituted(unittest.TestCase):
    """Раньше при отсутствии данных за второй год подставлялось значение
    соседнего: в документе оба года выглядели одинаковыми цифрами, а в
    отчёт о незаполненных тегах такая ячейка не попадала."""

    def _doc_with_tags(self):
        doc = Document()
        table = doc.add_table(rows=1, cols=2)
        table.cell(0, 0).text = "{{OKVED_A_VneAkt_23}}"
        table.cell(0, 1).text = "{{OKVED_A_VneAkt_24}}"
        return doc, table

    def test_absent_year_becomes_a_dash_not_the_other_year(self):
        doc, table = self._doc_with_tags()
        master = {"A": {"VneAkt_23": "100"}}  # данных за 24 год нет

        unfilled = fill_word_template_by_tags_v2(doc, master)

        self.assertEqual(table.cell(0, 0).text, "100")
        self.assertEqual(table.cell(0, 1).text, "-",
                         "значение за соседний год подставлено вместо прочерка")

    def test_absent_year_is_reported_as_unfilled(self):
        doc, _ = self._doc_with_tags()
        master = {"A": {"VneAkt_23": "100"}}

        unfilled = fill_word_template_by_tags_v2(doc, master)

        tags = [row["Тег (для разработчика)"] for row in unfilled]
        self.assertIn("OKVED_A_VneAkt_24", tags,
                      "незаполненная ячейка не попала в отчёт")


class TestDashBoldNormalisation(unittest.TestCase):
    """В исходном документе встречаются строки, где числа набраны жирным, а
    прочерк в той же строке — нет."""

    def _row_with_odd_dash(self):
        doc = Document()
        table = doc.add_table(rows=1, cols=4)
        for idx, text in enumerate(["Водоснабжение", "307388", "277378", "-"]):
            run = table.cell(0, idx).paragraphs[0].add_run(text)
            run.bold = (idx != 3)
        return doc, table

    def test_lone_non_bold_dash_is_aligned_with_its_row(self):
        doc, table = self._row_with_odd_dash()
        fixed = normalize_dash_bold_in_document(doc)

        self.assertEqual(fixed, 1)
        self.assertTrue(table.cell(0, 3).paragraphs[0].runs[0].bold)

    def test_row_that_is_entirely_non_bold_is_left_alone(self):
        doc = Document()
        table = doc.add_table(rows=1, cols=3)
        for idx, text in enumerate(["Строка", "12", "-"]):
            table.cell(0, idx).paragraphs[0].add_run(text).bold = False

        self.assertEqual(normalize_dash_bold_in_document(doc), 0)
        self.assertFalse(table.cell(0, 2).paragraphs[0].runs[0].bold)

    def test_dominant_bold_excludes_the_cell_being_judged(self):
        _, table = self._row_with_odd_dash()
        row = table.rows[0]
        self.assertTrue(dominant_row_bold(row, exclude_cell=row.cells[3]))


class TestTableAutofit(unittest.TestCase):
    """Автоподбор ширины должен трогать только таблицы, которые реально
    вылезают за печатную область: у аккуратно выставленной вручную таблицы
    растягивание на всю ширину окна испортило бы вёрстку."""

    def _usable_twips(self, doc):
        sec = doc.sections[0]
        return int((sec.page_width - sec.left_margin - sec.right_margin - sec.gutter) / 635)

    def test_overflowing_table_is_switched_to_window_width(self):
        doc = Document()
        table = doc.add_table(rows=1, cols=2)
        for col in table.columns:
            col.width = Twips(self._usable_twips(doc))  # вдвое шире страницы
        table.autofit = False

        self.assertEqual(autofit_tables_to_window(doc), 1)
        self.assertTrue(table.autofit)

    def test_table_that_fits_is_left_untouched(self):
        doc = Document()
        table = doc.add_table(rows=1, cols=2)
        for col in table.columns:
            col.width = Twips(self._usable_twips(doc) // 4)
        table.autofit = False

        self.assertEqual(autofit_tables_to_window(doc), 0)
        self.assertFalse(table.autofit, "ширина аккуратной таблицы изменена без нужды")


if __name__ == "__main__":
    unittest.main()
